"""Plain-text extraction for uploaded deal documents (PDF / PPTX).

Turns an uploaded file's raw bytes into the ``raw_text`` a :class:`SourceItem`
needs so it can flow through the existing, unchanged pipeline. Extraction is
purely local (no outbound network access) and never invents content: pages or
slides that carry no extractable text simply contribute nothing.
"""

from __future__ import annotations

from io import BytesIO


class UnsupportedFileType(ValueError):
    """Raised when an uploaded file's extension/content-type isn't supported."""


_PDF_EXTENSIONS = {".pdf"}
_PPTX_EXTENSIONS = {".pptx"}
# Legacy binary .ppt is not a zip/XML container (python-pptx can't read it).
_UNSUPPORTED_EXTENSIONS = {".ppt"}


def _extension_of(filename: str) -> str:
    name = (filename or "").strip().lower()
    dot = name.rfind(".")
    return name[dot:] if dot != -1 else ""


def extract_text_from_pdf(data: bytes) -> str:
    """Extract text from a PDF's bytes, page by page."""

    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(p.strip() for p in pages if p.strip())


def extract_text_from_pptx(data: bytes) -> str:
    """Extract text from a PPTX's bytes, slide by slide (shapes + tables)."""

    from pptx import Presentation

    prs = Presentation(BytesIO(data))
    slides_text: list[str] = []
    for slide in prs.slides:
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text
                if text and text.strip():
                    parts.append(text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if parts:
            slides_text.append("\n".join(parts))
    return "\n\n".join(slides_text)


def extract_text(filename: str, data: bytes) -> str:
    """Dispatch to the right extractor based on ``filename``'s extension.

    Raises :class:`UnsupportedFileType` for anything other than ``.pdf`` /
    ``.pptx`` (including legacy ``.ppt``, which isn't a format python-pptx can
    parse).
    """

    ext = _extension_of(filename)
    if ext in _PDF_EXTENSIONS:
        return extract_text_from_pdf(data)
    if ext in _PPTX_EXTENSIONS:
        return extract_text_from_pptx(data)
    if ext in _UNSUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(
            "Legacy .ppt is not supported — please save/export as .pptx and re-upload."
        )
    raise UnsupportedFileType(f"Unsupported file type {ext or '(none)'!r} — upload a .pdf or .pptx.")
