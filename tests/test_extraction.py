"""Tests for app.extraction: PDF/PPTX -> plain text."""

from __future__ import annotations

import pytest

from app.extraction import (
    UnsupportedFileType,
    extract_text,
    extract_text_from_pdf,
    extract_text_from_pptx,
)


def _make_minimal_pdf(text: str) -> bytes:
    """Build a tiny, byte-exact single-page PDF containing ``text``."""

    content = f"BT /F1 24 Tf 10 100 Td ({text}) Tj ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> "
        "/MediaBox [0 0 200 200] /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode("latin-1")
        out += body.encode("latin-1")
        out += b"\nendobj\n"

    xref_offset = len(out)
    n = len(objs) + 1
    out += f"xref\n0 {n}\n".encode("latin-1")
    out += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += f"{off:010d} 00000 n \n".encode("latin-1")
    out += f"trailer\n<< /Size {n} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode("latin-1")
    return bytes(out)


def _make_minimal_pptx(*slide_texts: str) -> bytes:
    from io import BytesIO

    from pptx import Presentation

    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    for text in slide_texts:
        slide = prs.slides.add_slide(blank_layout)
        box = slide.shapes.add_textbox(0, 0, prs.slide_width, prs.slide_height)
        box.text_frame.text = text
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_extract_text_from_pdf_returns_page_text():
    pdf_bytes = _make_minimal_pdf("Acme Bookkeeping For Sale")
    text = extract_text_from_pdf(pdf_bytes)
    assert "Acme Bookkeeping For Sale" in text


def test_extract_text_from_pptx_collects_every_slide():
    pptx_bytes = _make_minimal_pptx("Slide one intro", "Asking price $2.5M EBITDA $600k")
    text = extract_text_from_pptx(pptx_bytes)
    assert "Slide one intro" in text
    assert "Asking price $2.5M EBITDA $600k" in text


def test_extract_text_dispatches_by_extension():
    pdf_bytes = _make_minimal_pdf("Dispatch check")
    assert "Dispatch check" in extract_text("teaser.pdf", pdf_bytes)

    pptx_bytes = _make_minimal_pptx("Deck check")
    assert "Deck check" in extract_text("teaser.PPTX", pptx_bytes)


def test_extract_text_rejects_legacy_ppt():
    with pytest.raises(UnsupportedFileType):
        extract_text("teaser.ppt", b"not really parseable")


def test_extract_text_rejects_unknown_extension():
    with pytest.raises(UnsupportedFileType):
        extract_text("notes.txt", b"plain text")
